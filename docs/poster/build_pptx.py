#!/usr/bin/env python3
"""전시 포스터 PPTX 생성 — 의료 Physical AI 솔루션 경진대회 템플릿(Physical AI 빈 양식)에 내용 배치.

사용: python3 build_pptx.py
입력: ASSETS 경로의 이미지들 + 대회 템플릿 pptx
출력: PhysicalAI_MediScribe_전시포스터.pptx
"""
import pathlib
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = pathlib.Path(__file__).parent
A = HERE / "assets"        # 포스터에 쓰는 이미지·도식·대회 템플릿
TEMPLATE = A / "competition_template.pptx"
OUT = HERE / "PhysicalAI_MediScribe_전시포스터.pptx"

MM = 36000                      # 1mm in EMU
def mm(v): return Emu(int(v * MM))
def pt(v_mm): return Pt(round(v_mm * 2.835, 1))   # mm 글자높이 → pt

DEEP = RGBColor(0x43, 0x17, 0x6F)
MID  = RGBColor(0x6B, 0x3C, 0x9E)
INK  = RGBColor(0x1C, 0x15, 0x26)
SUB  = RGBColor(0x5B, 0x51, 0x68)
LINE = RGBColor(0xD9, 0xD0, 0xE4)
TEAL = RGBColor(0x0E, 0x9A, 0x96)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Paperlogy"

# ---- 레이아웃 (mm) ----
X0, X1 = 49.4, 900.6
CW = X1 - X0                    # 851.2
COLW = (CW - 8) / 2             # 421.6
XL, XR = X0, X0 + COLW + 8

def _prep_crops():
    """원본 사진을 레이아웃 비율에 맞게 잘라 assets/ 에 둔다 (없을 때만)."""
    from PIL import Image
    jobs = [("map_poster.png", "map_pptx.png", 421.6 / 180, 0.40),
            ("photo_full.jpg", "photo_full_c.png", 208.8 / 132, 0.62),
            ("photo_open.jpg", "photo_open_c.png", 208.8 / 132, 0.50)]
    for src, dst, aspect, focus in jobs:
        out = A / dst
        if out.exists():
            continue
        im = Image.open(A / src).convert("RGB")
        w, h = im.size
        if w / h > aspect:
            nw = int(h * aspect); im = im.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
        else:
            nh = int(w / aspect); y = int((h - nh) * focus); im = im.crop((0, y, w, y + nh))
        im.save(out, quality=92)


_prep_crops()
prs = Presentation(str(TEMPLATE))

# Physical AI 빈 양식(6번째)만 남기고 삭제
RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
lst = prs.slides._sldIdLst
for i, s in enumerate(list(lst)):
    if i != 5:
        prs.part.drop_rel(s.get(RID))
        lst.remove(s)
slide = prs.slides[0]
S = slide.shapes


def tb(x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = S.add_textbox(mm(x), mm(y), mm(w), mm(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, first=False, size=8, color=INK, bold=False, align=PP_ALIGN.LEFT, space_after=0, line=1.25):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    p._size, p._color, p._bold = size, color, bold
    return p


def run(p, text, size=None, color=None, bold=None):
    r = p.add_run(); r.text = text
    f = r.font
    f.name = FONT
    f.size = pt(size if size is not None else p._size)
    f.color.rgb = color if color is not None else p._color
    f.bold = p._bold if bold is None else bold
    return r


def bar(x, y, w, text, h=14.5, fill=DEEP):
    sh = S.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mm(x), mm(y), mm(w), mm(h))
    sh.adjustments[0] = 0.12
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = False
    tf.margin_left = mm(5.5); tf.margin_right = mm(4)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = text
    r.font.name = FONT; r.font.size = pt(9.0); r.font.bold = True; r.font.color.rgb = WHITE
    return y + h


def pic(path, x, y, w, h):
    S.add_picture(str(A / path), mm(x), mm(y), mm(w), mm(h))
    return y + h


def rule(x, y, w, color=DEEP, weight=1.4):
    ln = S.add_shape(MSO_SHAPE.RECTANGLE, mm(x), mm(y), mm(w), mm(weight))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


# ============ 헤더 ============
tf = tb(X0, 203, 520, 34)
p = para(tf, first=True, size=27, color=DEEP, bold=True, line=1.0)
run(p, "MediScribe")

tf = tb(X0, 240, 520, 14)
p = para(tf, first=True, size=10.5, color=MID, bold=True, line=1.1)
run(p, "회진에 동행하는 진료 요약 로봇")

tf = tb(X1 - 440, 202, 440, 52)
p = para(tf, first=True, size=8.6, color=INK, bold=True, align=PP_ALIGN.RIGHT, line=1.3)
run(p, "연세대학교")
p = para(tf, size=7.0, color=SUB, align=PP_ALIGN.RIGHT, line=1.4)
run(p, "건설환경공학과 ", bold=True, color=MID); run(p, "남윤상 · 김진찬")
p = para(tf, size=7.0, color=SUB, align=PP_ALIGN.RIGHT, line=1.4)
run(p, "미래캠퍼스 간호학과 ", bold=True, color=MID); run(p, "지민선 · 김해인")

rule(X0, 262, CW)

# ============ 리드 문장 ============
tf = tb(X0, 272, CW, 30)
p = para(tf, first=True, size=8.8, color=INK, line=1.35)
run(p, "의사–환자 대화를 듣고 ")
run(p, "환자가 가져갈 설명문", bold=True, color=DEEP)
run(p, "을 만들어, ")
run(p, "의사가 검수한 내용만", bold=True, color=DEEP)
run(p, " 그 자리에서 종이와 QR로 전달합니다. 음성인식·요약·주행이 모두 로봇 안에서 돌아가 ")
run(p, "대화는 병원 밖으로 나가지 않습니다.", bold=True, color=DEEP)

Y = 318

# ============ 왼쪽 단 ============
y = bar(XL, Y, COLW, "개발 배경 — 설명은 말로 한 번 끝난다")
tf = tb(XL, y + 5, COLW, 88)
for i, (a, b, c) in enumerate([
    ("· 회진 설명은 평균 수 분. 환자는 ", "약 용량·재검 일정·위험 신호", "를 기억하지 못한 채 병실을 나선다."),
    ("· 의사가 매번 요약을 적어주는 것은 진료 시간상 불가능하다.", "", ""),
    ("· 기존 음성기록 제품은 ", "의무기록용이고 클라우드로 전송", "된다. 환자용 설명문이 아니며, 대화가 외부로 나간다."),
]):
    p = para(tf, first=(i == 0), size=8.2, color=INK, line=1.32, space_after=6)
    run(p, a)
    if b: run(p, b, bold=True, color=DEEP)
    if c: run(p, c)
tf = tb(XL, 426, COLW, 14)
p = para(tf, first=True, size=7.0, color=SUB, line=1.3)
run(p, "→ 설명이 끝나는 그 자리에서, 환자가 손에 쥘 종이 한 장을 만든다.")

y = bar(XL, 448, COLW, "작동 방식 — 8단계 파이프라인")
y = pic("pipe.png", XL, y + 4, COLW, COLW * 900 / 2640)
tf = tb(XL, y + 5, COLW, 28)
p = para(tf, first=True, size=7.0, color=SUB, line=1.38)
run(p, "ReSpeaker 6-mic → Riva 한국어 스트리밍 인식 → 임상 내용 없는 발화 차단 → 소형 LLM 추출 → "
       "원문 대조 검증 → 규칙으로 용량·일정·경고 회수 → 의사 승인 → 출력. "
       "전 구간 Jetson Orin NX 16GB 온디바이스.")

y = bar(XL, 646, COLW, "자율주행 — 엔코더 없는 차량의 SLAM", fill=MID)
y = pic("map_pptx.png", XL, y + 4, COLW, 180)
tf = tb(XL, y + 4, COLW, 20)
p = para(tf, first=True, size=6.8, color=SUB, line=1.33)
run(p, "수동 주행 한 번으로 작성한 복도·교차로 점유격자 지도. 청록색은 주행 궤적. 바퀴 엔코더가 없어 ")
run(p, "매 스캔의 지도 정합 결과를 오도메트리로 사용", bold=True, color=DEEP)
run(p, "한다. ROS2 없이 순수 Python 구현.")

# ============ 오른쪽 단 ============
y = bar(XR, Y, COLW, "로봇 플랫폼")
pw = (COLW - 4) / 2
pic("photo_full_c.png", XR, y + 4, pw, 132)
pic("photo_open_c.png", XR + pw + 4, y + 4, pw, 132)
cy = y + 4 + 132 + 1.8
for i, (lab, desc) in enumerate([("외관", "상부 장비함과 전방 2D 라이다"), ("내부", "Jetson · 제어 보드 · 배터리")]):
    tf = tb(XR + i * (pw + 4), cy, pw, 9)
    p = para(tf, first=True, size=6.2, color=SUB, line=1.2)
    run(p, lab + "  ", bold=True, color=DEEP); run(p, desc)
tf = tb(XR, cy + 11, COLW, 24)
p = para(tf, first=True, size=7.0, color=SUB, line=1.38)
run(p, "1/5 스케일 애커만 조향 차량. 상부 장비함에 Jetson Orin NX 16GB, 전방 2D 라이다(STL-27L), "
       "6채널 마이크 어레이와 배터리를 싣고 하위 제어는 Arduino Mega가 맡는다.")

y = bar(XR, 508, COLW, "검증 결과")
rows = [("지표", "결과", "측정 조건"),
        ("요약 안전 정보 보존율", "98.1%", "모의 회진 14건 / 약 용량·일정·경고 항목 52→53 중"),
        ("요약 전체 보존율", "94.7%", "동일 데이터셋 71/75"),
        ("의학용어 인식 (메트포르민)", "3/3", "마이크 30cm, 단어 부스팅 20 적용"),
        ("미검수 출력 차단", "100%", "승인 전 출력·QR·인쇄 경로 전수 시험, 전부 403"),
        ("라이다 수신", "10Hz", "2,172점/회전, CRC 오류 0% (921,600bps 실측)")]
tbl_h = 110
gf = S.add_table(len(rows), 3, mm(XR), mm(y + 5), mm(COLW), mm(tbl_h))
tbl = gf.table
tbl.columns[0].width = mm(152); tbl.columns[1].width = mm(52); tbl.columns[2].width = mm(COLW - 204)
for r, row in enumerate(rows):
    tbl.rows[r].height = mm(tbl_h / len(rows))
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.background()
        cell.margin_left = mm(2.4); cell.margin_right = mm(2)
        cell.margin_top = cell.margin_bottom = mm(1.2)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        rr = p.add_run(); rr.text = val
        f = rr.font; f.name = FONT
        if r == 0:
            f.size = pt(6.2); f.bold = True; f.color.rgb = SUB
        elif c == 1:
            f.size = pt(7.6); f.bold = True; f.color.rgb = DEEP
        elif c == 2:
            f.size = pt(6.0); f.color.rgb = SUB
        else:
            f.size = pt(6.8); f.color.rgb = INK

CHW = 384
pic("chart.png", XR + (COLW - CHW) / 2, 640, CHW, CHW * 960 / 2640)

y = bar(XR, 786, COLW, "차별점", fill=MID)
tf = tb(XR, y + 5, COLW, 80)
for i, (head_, body_) in enumerate([
    ("로봇 안에서 끝난다.", " 외부 서버 전송이 없어 대화 원문이 병원을 벗어나지 않고, 인터넷이 없는 구역에서도 동작한다."),
    ("모델을 믿지 않는다.", " 원문에 없는 숫자는 코드가 버리고, 모델이 흘린 용량·일정·경고는 코드가 원문에서 직접 회수한다. 소형 모델의 환각을 구조로 막았다."),
    ("의사가 승인해야 나간다.", " 검수 전에는 인쇄·QR이 모두 차단된다. 사람이 책임지는 지점을 코드로 고정했다."),
]):
    p = para(tf, first=(i == 0), size=7.8, color=INK, line=1.32, space_after=6)
    run(p, f"{i+1}  ", bold=True, color=TEAL)
    run(p, head_, bold=True, color=DEEP)
    run(p, body_)

# ============ 전폭: 기대효과 ============
bar(X0, 880, CW, "활용 방안 및 기대효과")
tf = tb(X0, 899, CW, 32)
for i, (k, v) in enumerate([
    ("환자", "설명을 종이로 가져간다. 복약 순응도와 재방문 일정 준수가 올라간다."),
    ("의료진", "같은 설명의 반복과 퇴원 후 전화 문의가 줄고, 남는 기록은 의사가 승인한 내용뿐이다."),
    ("병원·확장", "외부 전송이 없어 개인정보 부담이 작고 인터넷 없는 구역에서도 쓴다. 투약 안내·퇴원 교육·검사 전 준비사항으로 확장한다."),
]):
    p = para(tf, first=(i == 0), size=7.4, color=INK, line=1.3, space_after=2)
    run(p, "· "); run(p, k, bold=True, color=DEEP); run(p, " — " + v)

# ============ 하단: 실제 생성 예시 ============
EY = 938
bar(X0, EY, CW, "실제 생성 예시 — 의사가 말하면, 환자는 이 종이를 받는다")
cy = EY + 19
ch = 120
cw1, cw2, cw3 = 300, 245, 300
gap = 3.1
cx2 = X0 + cw1 + gap
cx3 = cx2 + cw2 + gap


def card(x, y, w, h, accent=False):
    sh = S.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mm(x), mm(y), mm(w), mm(h))
    sh.adjustments[0] = 0.025
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0xFA, 0xF7, 0xFD) if accent else WHITE
    sh.line.color.rgb = DEEP if accent else LINE
    sh.line.width = Pt(2.2 if accent else 1.2)
    sh.shadow.inherit = False


card(X0, cy, cw1, ch)
card(cx2, cy, cw2, ch)
card(cx3, cy, cw3, ch, accent=True)

# 1) 발화
tf = tb(X0 + 6, cy + 4.5, cw1 - 12, 8)
p = para(tf, first=True, size=5.8, color=SUB, bold=True); run(p, "① 진료실에서 오간 말  —  의사 ▸ 환자")
tf = tb(X0 + 6, cy + 14, cw1 - 12, 86)
lines = [
    ("의사", [("“지난번 검사 결과부터 보겠습니다. 공복 혈당이 ", 0), ("백팔십", 1),
             ("이고 당화혈색소가 ", 0), ("팔점이 퍼센트", 1), ("입니다. 목표가 칠 퍼센트 아래라 조금 높아요.”", 0)]),
    ("의사", [("“메트포르민을 하루 오백에서 ", 0), ("천 밀리그램", 1),
             ("으로 올리겠습니다. ", 0), ("아침저녁으로 한 알씩", 1), (" 식사와 함께 드세요.”", 0)]),
    ("환자", [("“속이 불편하면 어떡하죠?”", 0)]),
    ("의사", [("“대개 일주일이면 괜찮아집니다. 그리고 저혈당이 올 수 있으니 ", 0),
             ("사탕을 꼭 가지고 다니세요", 1), (". 식은땀이 나거나 어지러우면 바로 드시고요.”", 0)]),
    ("의사", [("“", 0), ("이 주 뒤에", 1), (" 혈액검사 다시 하고 결과 보고 조정하겠습니다.”", 0)]),
]
for i, (who, segs) in enumerate(lines):
    p = para(tf, first=(i == 0), size=6.2, color=RGBColor(0x2A, 0x23, 0x35), line=1.26, space_after=3)
    run(p, who + "  ", bold=True, color=TEAL if who == "의사" else SUB)
    for t, em in segs:
        run(p, t, bold=bool(em), color=DEEP if em else None)
tf = tb(X0 + 6, cy + ch - 16, cw1 - 12, 13)
p = para(tf, first=True, size=5.4, color=SUB, line=1.28)
run(p, "잡담은 임상 게이트가 걸러 모델에 가지 않는다. 한글 수사는 숫자로 바꿔 원문과 대조한다 — 백팔십→180, 팔점이→8.2, 천→1000.")

# 2) 검수 화면
tf = tb(cx2 + 5, cy + 4.5, cw2 - 10, 8)
p = para(tf, first=True, size=5.8, color=SUB, bold=True); run(p, "② 의사 검수 — 실제 구동 캡처")
shot_h = 84
shot_w = shot_h * 732 / 586
pic("review_shot.png", cx2 + (cw2 - shot_w) / 2, cy + 14, shot_w, shot_h)
tf = tb(cx2 + 5, cy + ch - 19, cw2 - 10, 16)
p = para(tf, first=True, size=5.4, color=SUB, line=1.28)
run(p, "인식 오류가 남은 줄을 의사가 고치고 승인한다. ")
run(p, "승인 전에는 인쇄·QR 경로가 모두 403으로 차단된다.", bold=True, color=DEEP)

# 3) 환자 메모
tf = tb(cx3 + 6, cy + 4.5, cw3 - 12, 8)
p = para(tf, first=True, size=5.8, color=DEEP, bold=True); run(p, "③ 환자가 받아 가는 메모  —  A6 인쇄 · QR")
tf = tb(cx3 + 6, cy + 14, cw3 - 12, 86)
memo = [("오늘 검사 결과", ["공복 혈당 180", "당화혈색소 8.2% (목표 7% 미만)"]),
        ("처방 · 약", ["메트포르민 500mg → 1000mg 증량", "아침·저녁 1알씩, 식사와 함께"]),
        ("주의하세요", ["저혈당 대비해 사탕 휴대", "식은땀·어지럼 시 바로 섭취", "증량 초기 속쓰림 가능, 보통 1주 내 호전"]),
        ("다음 진료", ["2주 뒤 혈액검사 후 용량 조정"])]
first = True
for g, ls in memo:
    p = para(tf, first=first, size=5.8, color=DEEP, bold=True, line=1.2, space_after=1)
    run(p, g); first = False
    for l in ls:
        p = para(tf, size=6.2, color=INK, line=1.24, space_after=1)
        run(p, "   · " + l)
    tf.paragraphs[-1].space_after = Pt(4)
tf = tb(cx3 + 6, cy + ch - 13, cw3 - 12, 10)
p = para(tf, first=True, size=5.6, color=DEEP, bold=True, line=1.2)
run(p, "● 의사 검수 완료 — 승인본만 환자에게 나간다")

# ============ 꼬리 ============
rule(X0, 1080, CW, color=LINE, weight=0.6)
tf = tb(X0, 1085, CW, 10)
p = para(tf, first=True, size=5.8, color=SUB, line=1.2)
run(p, "온디바이스   ·   Human-in-the-loop   ·   ROS2 미사용 · 순수 Python        ", bold=True, color=DEEP)
run(p, "Jetson Orin NX 16GB · NVIDIA Riva ko-KR · LDROBOT STL-27L · Arduino Mega 2560")

prs.save(str(OUT))
print("saved:", OUT)
